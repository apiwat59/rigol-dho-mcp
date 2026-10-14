"""
MCP server for Rigol DHO900 series (DHO924S) - capture & report.

Tools are grouped as:
  - connection / status
  - acquisition control (run/stop/single/autoscale/settings)
  - capture (screenshot, measurements, waveform CSV + plot)
  - report session (start report, captures auto-attached, export DOCX/HTML)

Run:  python rigol_dho_mcp.py
Env:  RIGOL_RESOURCE  VISA resource string (default: auto-detect first Rigol)
      RIGOL_REPORT_DIR  output folder (default: ./reports next to this file)
"""
from __future__ import annotations

import base64
import csv
import datetime as dt
import json
import os
import re
import threading
from pathlib import Path
from typing import Annotated

import numpy as np
import pyvisa
from fastmcp import FastMCP
from fastmcp.utilities.types import Image
from pydantic import Field

BASE_DIR = Path(__file__).resolve().parent
REPORT_ROOT = Path(os.environ.get("RIGOL_REPORT_DIR", BASE_DIR / "reports"))
RESOURCE = os.environ.get("RIGOL_RESOURCE", "")

INVALID = 9.0e37  # Rigol returns 9.9E37 for "no valid measurement"

DEFAULT_MEAS = ["VPP", "VMAX", "VMIN", "VAVG", "VRMS", "FREQ", "PER",
                "RTIM", "FTIM", "PDUT"]
MEAS_UNITS = {
    "VPP": "V", "VMAX": "V", "VMIN": "V", "VAVG": "V", "VRMS": "V", "VTOP": "V",
    "VBAS": "V", "VAMP": "V", "VUPP": "V", "VMID": "V", "VLOW": "V", "PVRM": "V",
    "FREQ": "Hz", "PER": "s", "RTIM": "s", "FTIM": "s", "PWID": "s", "NWID": "s",
    "TVMAX": "s", "TVMIN": "s", "PDUT": "%", "NDUT": "%", "OVER": "%", "PRES": "%",
    "PSL": "V/s", "NSL": "V/s", "MAR": "V*s", "MPAR": "V*s",
}

mcp = FastMCP(
    "rigol-dho",
    instructions=(
        "Controls a Rigol DHO924S oscilloscope over VISA for capturing "
        "screenshots, measurements and waveforms for test reports. "
        "Typical flow: report_start -> (configure / single) -> capture_screenshot "
        "+ get_measurements per test point -> report_export. "
        "Channels are 1-4; sources are CHAN1..CHAN4."
    ),
)


# --------------------------------------------------------------------------- #
# Instrument connection
# --------------------------------------------------------------------------- #
class Scope:
    def __init__(self) -> None:
        self.lock = threading.RLock()
        self.rm: pyvisa.ResourceManager | None = None
        self.inst = None
        self.idn = ""

    def _find_resource(self) -> str:
        if RESOURCE:
            return RESOURCE
        for r in self.rm.list_resources():
            if r.upper().startswith("USB") and "0X1AB1" in r.upper():
                return r
        for r in self.rm.list_resources():
            if r.startswith("TCPIP"):
                try:
                    i = self.rm.open_resource(r, open_timeout=1500)
                    i.timeout = 1500
                    idn = i.query("*IDN?")
                    i.close()
                    if "RIGOL" in idn and "DHO" in idn:
                        return r
                except Exception:
                    pass
        raise RuntimeError("No Rigol DHO found. Set RIGOL_RESOURCE or check USB/LAN.")

    def connect(self, resource: str | None = None):
        with self.lock:
            self.close()
            if self.rm is None:
                self.rm = pyvisa.ResourceManager()
            res = resource or self._find_resource()
            self.inst = self.rm.open_resource(res)
            self.inst.timeout = 20000
            self.inst.chunk_size = 1024 * 1024
            self.idn = self.inst.query("*IDN?").strip()
            self.resource = res
            return self.idn

    def close(self):
        if self.inst is not None:
            try:
                self.inst.close()
            except Exception:
                pass
        self.inst = None

    def _ensure(self):
        if self.inst is None:
            self.connect()

    def _retry(self, fn):
        with self.lock:
            self._ensure()
            try:
                return fn()
            except (pyvisa.errors.VisaIOError, pyvisa.errors.InvalidSession):
                self.connect(getattr(self, "resource", None))
                return fn()

    def write(self, cmd: str):
        return self._retry(lambda: self.inst.write(cmd))

    def query(self, cmd: str) -> str:
        return self._retry(lambda: self.inst.query(cmd).strip())

    def sync(self):
        """Block until the scope has processed all previous commands (it drops/defers
        settings sent back-to-back without this)."""
        return self.query("*OPC?")

    def set(self, cmd: str):
        """Write a setting command and wait for it to be applied."""
        self.write(cmd)
        self.sync()

    def query_block(self, cmd: str) -> bytes:
        """Query returning IEEE 488.2 definite-length block (#9NNNNNNNNN...)."""
        def run():
            self.inst.write(cmd)
            raw = self.inst.read_raw()
            if not raw.startswith(b"#"):
                raise RuntimeError(f"Unexpected block header: {raw[:20]!r}")
            nd = int(raw[1:2])
            n = int(raw[2:2 + nd])
            data = raw[2 + nd:]
            while len(data) < n:
                data += self.inst.read_raw()
            return data[:n]
        return self._retry(run)

    def qfloat(self, cmd: str) -> float | None:
        try:
            v = float(self.query(cmd))
        except ValueError:
            return None
        return None if abs(v) >= INVALID else v


scope = Scope()


# --------------------------------------------------------------------------- #
# Report session state
# --------------------------------------------------------------------------- #
class Report:
    def __init__(self) -> None:
        self.dir: Path | None = None
        self.meta: dict = {}

    @property
    def active(self) -> bool:
        return self.dir is not None

    def manifest(self) -> Path:
        return self.dir / "report.json"

    def save(self):
        self.manifest().write_text(json.dumps(self.meta, indent=2, ensure_ascii=False),
                                   encoding="utf-8")

    def start(self, title: str, **info) -> Path:
        stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        slug = re.sub(r"[^\w\-]+", "_", title, flags=re.UNICODE).strip("_")[:40] or "report"
        self.dir = REPORT_ROOT / f"{stamp}_{slug}"
        self.dir.mkdir(parents=True, exist_ok=True)
        self.meta = {"title": title, "created": dt.datetime.now().isoformat(timespec="seconds"),
                     "instrument": scope.idn, **info, "items": []}
        self.save()
        return self.dir

    def load(self, folder: Path):
        self.dir = folder
        self.meta = json.loads((folder / "report.json").read_text(encoding="utf-8"))

    def out_dir(self) -> Path:
        if self.dir is not None:
            return self.dir
        d = REPORT_ROOT / "_loose"
        d.mkdir(parents=True, exist_ok=True)
        return d

    def next_name(self, label: str, ext: str) -> Path:
        n = len(self.meta.get("items", [])) + 1 if self.active else 0
        slug = re.sub(r"[^\w\-]+", "_", label or "capture", flags=re.UNICODE).strip("_")[:40]
        stamp = dt.datetime.now().strftime("%H%M%S")
        prefix = f"{n:02d}_" if self.active else dt.datetime.now().strftime("%Y%m%d_")
        return self.out_dir() / f"{prefix}{slug}_{stamp}.{ext}"

    def add(self, item: dict):
        if not self.active:
            return
        item.setdefault("time", dt.datetime.now().isoformat(timespec="seconds"))
        self.meta["items"].append(item)
        self.save()


report = Report()


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def _src(ch: int | str) -> str:
    s = str(ch).upper()
    if s.isdigit():
        return f"CHAN{s}"
    if s.startswith("CH") and s[-1].isdigit() and len(s) <= 5:
        return f"CHAN{s[-1]}"
    return s


def fmt_eng(v: float | None, unit: str = "") -> str:
    if v is None:
        return "----"
    if unit == "%":  # scope returns duty/overshoot as a ratio (0.5 = 50 %)
        return f"{v * 100:.2f} %"
    if v == 0:
        return f"0 {unit}".strip()
    prefixes = [(1e9, "G"), (1e6, "M"), (1e3, "k"), (1, ""), (1e-3, "m"),
                (1e-6, "µ"), (1e-9, "n"), (1e-12, "p")]
    for f, p in prefixes:
        if abs(v) >= f:
            return f"{v / f:.4g} {p}{unit}".strip()
    return f"{v:.4g} {unit}".strip()


def _settings_snapshot() -> dict:
    snap = {
        "timebase_s_div": scope.qfloat(":TIM:MAIN:SCAL?"),
        "timebase_offset_s": scope.qfloat(":TIM:MAIN:OFFS?"),
        "sample_rate_Sa_s": scope.qfloat(":ACQ:SRAT?"),
        "memory_depth": scope.query(":ACQ:MDEP?"),
        "acq_type": scope.query(":ACQ:TYPE?"),
        "trigger": {
            "mode": scope.query(":TRIG:MODE?"),
            "sweep": scope.query(":TRIG:SWE?"),
            "status": scope.query(":TRIG:STAT?"),
        },
        "channels": {},
    }
    if snap["trigger"]["mode"] == "EDGE":
        snap["trigger"].update({
            "source": scope.query(":TRIG:EDGE:SOUR?"),
            "slope": scope.query(":TRIG:EDGE:SLOP?"),
            "level_V": scope.qfloat(":TRIG:EDGE:LEV?"),
        })
    for ch in range(1, 5):
        if scope.query(f":CHAN{ch}:DISP?") == "1":
            snap["channels"][f"CH{ch}"] = {
                "scale_V_div": scope.qfloat(f":CHAN{ch}:SCAL?"),
                "offset_V": scope.qfloat(f":CHAN{ch}:OFFS?"),
                "coupling": scope.query(f":CHAN{ch}:COUP?"),
                "probe": scope.qfloat(f":CHAN{ch}:PROB?"),
                "bw_limit": scope.query(f":CHAN{ch}:BWL?"),
            }
    try:
        snap["generator"] = _gen_status()
    except Exception:
        pass
    return snap


def _measure(source: str, items: list[str]) -> dict:
    out = {}
    for it in items:
        it = it.upper()
        out[it] = scope.qfloat(f":MEAS:ITEM? {it},{source}")
    return out


def _read_waveform(source: str, mode: str, max_points: int):
    """Return (t, v, preamble dict). mode NORM (screen) or RAW (memory, needs STOP)."""
    scope.set(f":WAV:SOUR {source}")
    scope.set(f":WAV:MODE {mode}")
    scope.set(":WAV:FORM WORD")
    if mode.upper() == "RAW":
        # Without this the preamble/readout stays at the previous point count (e.g. 10k).
        mdep = scope.qfloat(":ACQ:MDEP?")
        if mdep:
            scope.set(f":WAV:POIN {int(mdep)}")
    pre = scope.query(":WAV:PRE?").split(",")
    fmt, typ, npts, cnt, xinc, xorg, xref, yinc, yorg, yref = pre
    npts, xinc, xorg, xref = int(float(npts)), float(xinc), float(xorg), float(xref)
    yinc, yorg, yref = float(yinc), float(yorg), float(yref)
    total = min(npts, max_points) if max_points > 0 else npts
    chunk = 250_000
    raw = bytearray()
    start = 1
    while start <= total:
        stop = min(start + chunk - 1, total)
        scope.write(f":WAV:STAR {start}")
        scope.set(f":WAV:STOP {stop}")
        raw += scope.query_block(":WAV:DATA?")
        start = stop + 1
    scope.write(":WAV:STAR 1")
    codes = np.frombuffer(bytes(raw), dtype="<u2").astype(np.float64)
    v = (codes - yorg - yref) * yinc
    t = (np.arange(len(v)) - xref) * xinc + xorg
    info = {"points": len(v), "x_increment_s": xinc, "x_origin_s": xorg,
            "y_increment_V": yinc, "mode": mode}
    return t, v, info


def _plot(traces: dict, path: Path, title: str):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    colors = {"CHAN1": "#d4b000", "CHAN2": "#00b7eb", "CHAN3": "#d63fb5", "CHAN4": "#2a6fdb"}
    fig, ax = plt.subplots(figsize=(10, 4.5), dpi=130)
    for src, (t, v) in traces.items():
        span = t[-1] - t[0] if len(t) > 1 else 1
        scale, unit = next(((f, u) for f, u in [(1, "s"), (1e-3, "ms"), (1e-6, "µs"),
                            (1e-9, "ns")] if span >= f * 2), (1e-9, "ns"))
        step = max(1, len(t) // 200_000)
        ax.plot(t[::step] / scale, v[::step], lw=0.8, label=src.replace("CHAN", "CH"),
                color=colors.get(src))
        ax.set_xlabel(f"Time ({unit})")
    ax.set_ylabel("Voltage (V)")
    ax.set_title(title)
    ax.grid(True, alpha=0.3)
    ax.legend(loc="upper right")
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


# --------------------------------------------------------------------------- #
# Tools: connection / status
# --------------------------------------------------------------------------- #
@mcp.tool
def scope_connect(resource: Annotated[str, Field(description="VISA resource, blank = auto-detect")] = "") -> dict:
    """Connect (or reconnect) to the oscilloscope and return *IDN?."""
    idn = scope.connect(resource or None)
    return {"resource": scope.resource, "idn": idn}


@mcp.tool
def scope_list_resources() -> list[str]:
    """List VISA resources visible to this PC."""
    rm = scope.rm or pyvisa.ResourceManager()
    scope.rm = rm
    return list(rm.list_resources())


@mcp.tool
def scope_status() -> dict:
    """Current scope setup: timebase, sample rate, memory depth, trigger, enabled channels."""
    with scope.lock:
        scope._ensure()
        return {"idn": scope.idn, **_settings_snapshot()}


# --------------------------------------------------------------------------- #
# Tools: acquisition control
# --------------------------------------------------------------------------- #
@mcp.tool
def scope_control(action: Annotated[str, Field(
        description="run | stop | single | autoscale | force_trigger | clear")]) -> str:
    """Run/Stop/Single/Autoscale/Force trigger/Clear display."""
    cmd = {"run": ":RUN", "stop": ":STOP", "single": ":SING", "autoscale": ":AUT",
           "force_trigger": ":TFOR", "clear": ":CLE"}.get(action.lower())
    if not cmd:
        raise ValueError("action must be run|stop|single|autoscale|force_trigger|clear")
    scope.write(cmd)
    return f"{action} sent; trigger status={scope.query(':TRIG:STAT?')}"


@mcp.tool
def scope_wait_trigger(timeout_s: float = 10.0) -> str:
    """After 'single', wait until the acquisition has stopped (triggered). Returns final status."""
    import time
    t0 = time.time()
    st = ""
    while time.time() - t0 < timeout_s:
        st = scope.query(":TRIG:STAT?")
        if st == "STOP":
            return "STOP (captured)"
        time.sleep(0.2)
    return f"timeout, status={st}"


@mcp.tool
def configure_channel(
    channel: Annotated[int, Field(ge=1, le=4)],
    display: bool | None = None,
    scale_v_div: float | None = None,
    offset_v: float | None = None,
    coupling: Annotated[str | None, Field(description="DC | AC | GND")] = None,
    probe: Annotated[float | None, Field(description="Probe ratio e.g. 1, 10")] = None,
    bw_limit: Annotated[str | None, Field(description="OFF | 20M")] = None,
    label: str | None = None,
) -> dict:
    """Set channel parameters (only those given). Probe is applied first so scale/offset
    are interpreted at the new ratio; an edge-trigger level on this channel is kept."""
    c = f":CHAN{channel}"
    if display is not None:
        scope.set(f"{c}:DISP {'ON' if display else 'OFF'}")
    # The scope rescales the trigger level with the probe ratio; restore it once the
    # new scale/offset are in place (it is clamped to the visible range).
    trig_lev = None
    if probe is not None:
        if (scope.query(":TRIG:MODE?") == "EDGE"
                and scope.query(":TRIG:EDGE:SOUR?") == f"CHAN{channel}"):
            trig_lev = scope.qfloat(":TRIG:EDGE:LEV?")
        scope.set(f"{c}:PROB {probe:g}")
    if coupling:
        scope.set(f"{c}:COUP {coupling.upper()}")
    if bw_limit:
        scope.set(f"{c}:BWL {bw_limit.upper()}")
    if scale_v_div is not None:
        scope.set(f"{c}:SCAL {scale_v_div}")
    if offset_v is not None:
        scope.set(f"{c}:OFFS {offset_v}")
    if trig_lev is not None:
        scope.set(f":TRIG:EDGE:LEV {trig_lev}")
    if label is not None:
        scope.set(f'{c}:LAB:CONT "{label}"')
        scope.set(f"{c}:LAB:SHOW ON")
    return {"error": scope.query(":SYST:ERR?"),
            "scale": scope.qfloat(f"{c}:SCAL?"), "offset": scope.qfloat(f"{c}:OFFS?"),
            "coupling": scope.query(f"{c}:COUP?"), "probe": scope.qfloat(f"{c}:PROB?")}


@mcp.tool
def configure_timebase(scale_s_div: float | None = None, offset_s: float | None = None,
                       memory_depth: Annotated[str | None, Field(
                           description="AUTO | 1k | 10k | 100k | 1M | 10M | 25M | 50M")] = None,
                       acq_type: Annotated[str | None, Field(
                           description="NORM | PEAK | AVER | ULTR")] = None) -> dict:
    """Set horizontal scale/offset, memory depth, acquisition type."""
    if scale_s_div is not None:
        scope.set(f":TIM:MAIN:SCAL {scale_s_div}")
    if offset_s is not None:
        scope.set(f":TIM:MAIN:OFFS {offset_s}")
    if memory_depth:
        scope.set(f":ACQ:MDEP {memory_depth.upper()}")
    if acq_type:
        scope.set(f":ACQ:TYPE {acq_type.upper()}")
    return {"error": scope.query(":SYST:ERR?"),
            "scale": scope.qfloat(":TIM:MAIN:SCAL?"), "offset": scope.qfloat(":TIM:MAIN:OFFS?"),
            "srate": scope.qfloat(":ACQ:SRAT?"), "mdepth": scope.query(":ACQ:MDEP?")}


@mcp.tool
def configure_trigger_edge(source: Annotated[str, Field(description="CHAN1..CHAN4 or 1..4")] = "CHAN1",
                           level_v: float | None = None,
                           slope: Annotated[str, Field(description="POS | NEG | RFAL")] = "POS",
                           sweep: Annotated[str, Field(description="AUTO | NORM | SING")] = "AUTO") -> dict:
    """Configure an edge trigger."""
    s = _src(source)
    scope.set(":TRIG:MODE EDGE")
    scope.set(f":TRIG:EDGE:SOUR {s}")
    scope.set(f":TRIG:EDGE:SLOP {slope.upper()}")
    if level_v is not None:
        scope.set(f":TRIG:EDGE:LEV {level_v}")
    scope.set(f":TRIG:SWE {sweep.upper()}")
    return {"error": scope.query(":SYST:ERR?"), "source": scope.query(":TRIG:EDGE:SOUR?"),
            "level": scope.qfloat(":TRIG:EDGE:LEV?"), "sweep": scope.query(":TRIG:SWE?")}


@mcp.tool
def show_measurement_on_screen(item: Annotated[str, Field(description="e.g. VPP, FREQ, RTIM")],
                               source: str = "CHAN1", clear_first: bool = False) -> str:
    """Add a measurement to the scope's on-screen measurement panel (visible in screenshots)."""
    if clear_first:
        scope.set(":MEAS:CLE")  # DHO firmware rejects ":MEAS:CLE ALL" (-108)
    scope.set(f":MEAS:ITEM {item.upper()},{_src(source)}")
    return scope.query(":SYST:ERR?")


@mcp.tool
def scpi(command: str,
         timeout_ms: Annotated[int, Field(description="Query timeout")] = 3000) -> str:
    """Send a raw SCPI command. If it ends with '?', returns the response; otherwise returns SYST:ERR?.
    An unknown query times out quickly and returns the scope's error instead of hanging."""
    with scope.lock:
        scope._ensure()
        if not command.strip().endswith("?"):
            scope.set(command)
            return scope.query(":SYST:ERR?")
        old = scope.inst.timeout
        scope.inst.timeout = timeout_ms
        try:
            return scope.inst.query(command).strip()
        except pyvisa.errors.VisaIOError:
            scope.inst.clear()
            return f"no response (timeout); SYST:ERR? = {scope.inst.query(':SYST:ERR?').strip()}"
        finally:
            scope.inst.timeout = old


# --------------------------------------------------------------------------- #
# Tools: built-in signal generator (DHO924S "G")
# --------------------------------------------------------------------------- #
def _gen_status() -> dict:
    return {
        "output": scope.query(":SOUR:OUTP:STAT?") in ("1", "ON"),
        "function": scope.query(":SOUR:FUNC?"),
        "frequency_Hz": scope.qfloat(":SOUR:FREQ?"),
        "amplitude_Vpp": scope.qfloat(":SOUR:VOLT:AMPL?"),
        "offset_V": scope.qfloat(":SOUR:VOLT:OFFS?"),
        "square_duty_pct": scope.qfloat(":SOUR:FUNC:SQU:DUTY?"),
        "modulation": scope.query(":SOUR:MOD:STAT?") in ("1", "ON"),
    }


@mcp.tool
def generator_status() -> dict:
    """Read the built-in signal generator settings and whether its output is ON."""
    return _gen_status()


@mcp.tool
def configure_generator(
    output: bool | None = None,
    function: Annotated[str | None, Field(
        description="SIN | SQU | RAMP | PULS | DC | NOIS | ARB")] = None,
    frequency_hz: float | None = None,
    amplitude_vpp: float | None = None,
    offset_v: float | None = None,
    square_duty_pct: float | None = None,
) -> dict:
    """Set the built-in signal generator (only parameters given). Output is switched last."""
    if function:
        scope.set(f":SOUR:FUNC {function.upper()}")
    if frequency_hz is not None:
        scope.set(f":SOUR:FREQ {frequency_hz}")
    if amplitude_vpp is not None:
        scope.set(f":SOUR:VOLT:AMPL {amplitude_vpp}")
    if offset_v is not None:
        scope.set(f":SOUR:VOLT:OFFS {offset_v}")
    if square_duty_pct is not None:
        scope.set(f":SOUR:FUNC:SQU:DUTY {square_duty_pct}")
    if output is not None:
        scope.set(f":SOUR:OUTP:STAT {'ON' if output else 'OFF'}")
    return {"error": scope.query(":SYST:ERR?"), **_gen_status()}


# --------------------------------------------------------------------------- #
# Tools: capture
# --------------------------------------------------------------------------- #
@mcp.tool
def capture_screenshot(
    label: Annotated[str, Field(description="Short name e.g. 'TP3 3V3 ripple'")] = "screenshot",
    note: Annotated[str, Field(description="Description/observation for the report")] = "",
    measure_source: Annotated[str, Field(description="Also record measurements for this source (e.g. CHAN1). Blank = skip")] = "",
    measure_items: Annotated[list[str] | None, Field(description="Measurement items; default common set")] = None,
    return_image: bool = True,
):
    """Capture the scope screen as PNG (saved to the report folder). Optionally records
    measurements + setup alongside it. Returns the image so it can be inspected."""
    with scope.lock:
        png = scope.query_block(":DISP:DATA? PNG")
        path = report.next_name(label, "png")
        path.write_bytes(png)
        meas = {}
        if measure_source:
            for s in [x.strip() for x in measure_source.split(",") if x.strip()]:
                meas[_src(s)] = _measure(_src(s), measure_items or DEFAULT_MEAS)
        settings = _settings_snapshot()
    report.add({"type": "screenshot", "label": label, "note": note, "file": path.name,
                "measurements": meas, "settings": settings})
    summary = {"saved": str(path), "report_active": report.active}
    if meas:
        summary["measurements"] = {s: {k: fmt_eng(v, MEAS_UNITS.get(k, "")) for k, v in m.items()}
                                   for s, m in meas.items()}
    if return_image:
        return [json.dumps(summary, ensure_ascii=False), Image(data=png, format="png")]
    return summary


@mcp.tool
def get_measurements(
    source: Annotated[str, Field(description="CHAN1..CHAN4, comma-separated for several")] = "CHAN1",
    items: Annotated[list[str] | None, Field(description=(
        "Items e.g. VPP VMAX VMIN VAVG VRMS VTOP VBAS VAMP OVER PRES FREQ PER RTIM FTIM "
        "PWID NWID PDUT NDUT PSL NSL; default common set"))] = None,
    label: str = "",
    note: str = "",
    add_to_report: bool = False,
) -> dict:
    """Read measurements (values in SI units; None = not measurable)."""
    res = {}
    with scope.lock:
        for s in [x.strip() for x in source.split(",") if x.strip()]:
            res[_src(s)] = _measure(_src(s), items or DEFAULT_MEAS)
    if add_to_report:
        report.add({"type": "measurements", "label": label or "Measurements", "note": note,
                    "measurements": res})
    return {"raw": res, "formatted": {s: {k: fmt_eng(v, MEAS_UNITS.get(k, "")) for k, v in m.items()}
                                      for s, m in res.items()}}


@mcp.tool
def capture_waveform(
    sources: Annotated[str, Field(description="e.g. 'CHAN1' or 'CHAN1,CHAN2'")] = "CHAN1",
    label: str = "waveform",
    note: str = "",
    mode: Annotated[str, Field(description="NORM = screen data (fast); RAW = full memory (scope is stopped)")] = "NORM",
    max_points: Annotated[int, Field(description="Limit points for RAW (0 = all)")] = 1_000_000,
    make_plot: bool = True,
) -> dict:
    """Download waveform data -> CSV (time, volts per channel) + optional PNG plot in the report folder."""
    mode = mode.upper()
    srcs = [_src(s) for s in sources.split(",") if s.strip()]
    traces = {}
    infos = {}
    with scope.lock:
        if mode == "RAW":
            import time
            scope.write(":STOP")
            time.sleep(0.3)
        for s in srcs:
            t, v, info = _read_waveform(s, mode, max_points)
            traces[s] = (t, v)
            infos[s] = info
    csv_path = report.next_name(label, "csv")
    n = min(len(t) for t, _ in traces.values())
    t0 = next(iter(traces.values()))[0][:n]
    with csv_path.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["time_s"] + [f"{s}_V" for s in srcs])
        cols = [traces[s][1][:n] for s in srcs]
        for i in range(n):
            w.writerow([f"{t0[i]:.9e}"] + [f"{c[i]:.6e}" for c in cols])
    stats = {s: {"min_V": float(v.min()), "max_V": float(v.max()), "mean_V": float(v.mean()),
                 "rms_V": float(np.sqrt(np.mean(v ** 2))), "pp_V": float(v.max() - v.min())}
             for s, (t, v) in traces.items()}
    plot_path = None
    if make_plot:
        plot_path = csv_path.with_suffix(".plot.png")
        _plot(traces, plot_path, label)
    report.add({"type": "waveform", "label": label, "note": note, "file": csv_path.name,
                "plot": plot_path.name if plot_path else None, "info": infos, "stats": stats})
    return {"csv": str(csv_path), "plot": str(plot_path) if plot_path else None,
            "info": infos, "stats": stats}


# --------------------------------------------------------------------------- #
# Tools: report
# --------------------------------------------------------------------------- #
@mcp.tool
def report_start(title: str,
                 dut: Annotated[str, Field(description="Device under test / board / revision")] = "",
                 engineer: str = "",
                 project: str = "",
                 description: str = "") -> dict:
    """Start a new report session. Subsequent captures are saved into its folder and listed in the report."""
    with scope.lock:
        scope._ensure()
    d = report.start(title, dut=dut, engineer=engineer, project=project, description=description)
    return {"folder": str(d), "meta": report.meta}


@mcp.tool
def report_open(folder: Annotated[str, Field(description="Existing report folder name or path")]) -> dict:
    """Resume an existing report session."""
    p = Path(folder)
    if not p.is_absolute():
        p = REPORT_ROOT / folder
    report.load(p)
    return {"folder": str(p), "items": len(report.meta["items"])}


@mcp.tool
def report_list() -> list[str]:
    """List report folders."""
    if not REPORT_ROOT.exists():
        return []
    return sorted([p.name for p in REPORT_ROOT.iterdir() if (p / "report.json").exists()], reverse=True)


@mcp.tool
def report_status() -> dict:
    """Show the active report and its items."""
    if not report.active:
        return {"active": False}
    return {"active": True, "folder": str(report.dir),
            "items": [{"index": i + 1, "type": it["type"], "label": it.get("label"),
                       "note": it.get("note"), "result": it.get("result")}
                      for i, it in enumerate(report.meta["items"])]}


@mcp.tool
def report_add_note(text: str, heading: str = "",
                    result: Annotated[str, Field(description="Optional verdict: PASS | FAIL | INFO")] = "") -> str:
    """Add a free-text section (observation, conclusion, setup description) to the report."""
    if not report.active:
        raise RuntimeError("No active report. Call report_start first.")
    report.add({"type": "note", "label": heading, "note": text, "result": result.upper() or None})
    return f"added item #{len(report.meta['items'])}"


@mcp.tool
def report_edit_item(index: Annotated[int, Field(description="1-based item index from report_status")],
                     label: str | None = None, note: str | None = None,
                     result: Annotated[str | None, Field(description="PASS | FAIL | INFO | '' to clear")] = None,
                     limit: Annotated[str | None, Field(description="Spec/limit text e.g. 'Vpp < 50 mV'")] = None,
                     delete: bool = False) -> dict:
    """Edit label/note/result/limit of a report item, or delete it."""
    if not report.active:
        raise RuntimeError("No active report.")
    items = report.meta["items"]
    it = items[index - 1]
    if delete:
        items.pop(index - 1)
    else:
        if label is not None:
            it["label"] = label
        if note is not None:
            it["note"] = note
        if result is not None:
            it["result"] = result.upper() or None
        if limit is not None:
            it["limit"] = limit
    report.save()
    return {"items": len(items)}


@mcp.tool
def report_set_info(title: str | None = None, dut: str | None = None, engineer: str | None = None,
                    project: str | None = None, description: str | None = None,
                    conclusion: str | None = None) -> dict:
    """Update report header fields / overall conclusion."""
    if not report.active:
        raise RuntimeError("No active report.")
    for k, v in dict(title=title, dut=dut, engineer=engineer, project=project,
                     description=description, conclusion=conclusion).items():
        if v is not None:
            report.meta[k] = v
    report.save()
    return report.meta | {"items": len(report.meta["items"])}


def _meas_rows(meas: dict):
    for src, m in meas.items():
        for k, v in m.items():
            yield src.replace("CHAN", "CH"), k, fmt_eng(v, MEAS_UNITS.get(k, ""))


def _settings_line(s: dict) -> str:
    if not s:
        return ""
    parts = [f"Time {fmt_eng(s.get('timebase_s_div'), 's')}/div",
             f"{fmt_eng(s.get('sample_rate_Sa_s'), 'Sa/s')}", f"Mem {s.get('memory_depth')}"]
    tr = s.get("trigger", {})
    if tr.get("source"):
        parts.append(f"Trig {tr['source']} {tr.get('slope')} {fmt_eng(tr.get('level_V'), 'V')}")
    for ch, c in s.get("channels", {}).items():
        parts.append(f"{ch} {fmt_eng(c['scale_V_div'], 'V')}/div {c['coupling']} x{c['probe']:g}")
    g = s.get("generator")
    if g:
        parts.append(f"Gen {g['function']} {fmt_eng(g['frequency_Hz'], 'Hz')} "
                     f"{fmt_eng(g['amplitude_Vpp'], 'Vpp')} {'ON' if g['output'] else 'OFF'}")
    return " | ".join(parts)


def _export_docx(path: Path):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    m = report.meta
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(1.8)

    doc.add_heading(m.get("title", "Test Report"), 0)
    info = doc.add_table(rows=0, cols=2)
    info.style = "Light Grid Accent 1"
    for k, lab in [("project", "Project"), ("dut", "DUT"), ("engineer", "Engineer"),
                   ("created", "Date"), ("instrument", "Instrument")]:
        if m.get(k):
            r = info.add_row().cells
            r[0].text, r[1].text = lab, str(m[k])
    if m.get("description"):
        doc.add_paragraph()
        doc.add_paragraph(m["description"])

    verdicts = [it.get("result") for it in m["items"] if it.get("result")]
    if verdicts:
        doc.add_heading("Summary", 1)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ["#", "Test point", "Limit", "Result"]):
            c.text = h
        for i, it in enumerate(m["items"], 1):
            if it.get("result"):
                r = t.add_row().cells
                r[0].text, r[1].text = str(i), it.get("label") or ""
                r[2].text, r[3].text = it.get("limit") or "", it["result"]
                run = r[3].paragraphs[0].runs[0]
                run.bold = True
                run.font.color.rgb = RGBColor(0x1A, 0x7F, 0x37) if it["result"] == "PASS" else (
                    RGBColor(0xC0, 0x1C, 0x28) if it["result"] == "FAIL" else RGBColor(0x55, 0x55, 0x55))

    doc.add_heading("Results", 1)
    for i, it in enumerate(m["items"], 1):
        head = f"{i}. {it.get('label') or it['type'].title()}"
        if it.get("result"):
            head += f"  [{it['result']}]"
        doc.add_heading(head, 2)
        if it.get("limit"):
            doc.add_paragraph(f"Limit: {it['limit']}")
        if it.get("note"):
            doc.add_paragraph(it["note"])
        img = it.get("file") if it["type"] == "screenshot" else it.get("plot")
        if img and (report.dir / img).exists():
            doc.add_picture(str(report.dir / img), width=Cm(16))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if it.get("settings"):
            p = doc.add_paragraph(_settings_line(it["settings"]))
            p.runs[0].font.size = Pt(8)
            p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        if it.get("measurements"):
            rows = list(_meas_rows(it["measurements"]))
            t = doc.add_table(rows=1, cols=3)
            t.style = "Light List Accent 1"
            for c, h in zip(t.rows[0].cells, ["Source", "Parameter", "Value"]):
                c.text = h
            for row in rows:
                cells = t.add_row().cells
                for c, v in zip(cells, row):
                    c.text = v
        if it.get("stats"):
            t = doc.add_table(rows=1, cols=6)
            t.style = "Light List Accent 1"
            for c, h in zip(t.rows[0].cells, ["Source", "Min", "Max", "Mean", "RMS", "Pk-Pk"]):
                c.text = h
            for s, v in it["stats"].items():
                cells = t.add_row().cells
                vals = [s.replace("CHAN", "CH")] + [fmt_eng(v[k], "V") for k in
                                                    ["min_V", "max_V", "mean_V", "rms_V", "pp_V"]]
                for c, x in zip(cells, vals):
                    c.text = x
            doc.add_paragraph(f"Data file: {it['file']}").runs[0].font.size = Pt(8)
    if m.get("conclusion"):
        doc.add_heading("Conclusion", 1)
        doc.add_paragraph(m["conclusion"])
    doc.save(path)


def _export_html(path: Path):
    import html as H
    m = report.meta

    def img64(name):
        p = report.dir / name
        return "data:image/png;base64," + base64.b64encode(p.read_bytes()).decode() if p.exists() else ""

    def badge(r):
        if not r:
            return ""
        return f'<span class="b {r.lower()}">{H.escape(r)}</span>'

    out = [f"""<!doctype html><html><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{H.escape(m.get('title', 'Report'))}</title><style>
body{{font-family:Segoe UI,Arial,sans-serif;max-width:960px;margin:24px auto;padding:0 16px;color:#222}}
h1{{border-bottom:3px solid #1f5fa8;padding-bottom:6px}} h2{{margin-top:32px;color:#1f5fa8}}
table{{border-collapse:collapse;margin:8px 0}} td,th{{border:1px solid #ccc;padding:4px 10px;font-size:14px}}
th{{background:#eef3fa;text-align:left}} img{{max-width:100%;border:1px solid #ddd}}
.set{{color:#666;font-size:12px}} .b{{padding:2px 8px;border-radius:4px;color:#fff;font-size:13px}}
.pass{{background:#1a7f37}} .fail{{background:#c01c28}} .info{{background:#666}}
.item{{page-break-inside:avoid}}</style></head><body>
<h1>{H.escape(m.get('title', 'Report'))}</h1><table>"""]
    for k, lab in [("project", "Project"), ("dut", "DUT"), ("engineer", "Engineer"),
                   ("created", "Date"), ("instrument", "Instrument")]:
        if m.get(k):
            out.append(f"<tr><th>{lab}</th><td>{H.escape(str(m[k]))}</td></tr>")
    out.append("</table>")
    if m.get("description"):
        out.append(f"<p>{H.escape(m['description'])}</p>")
    if any(it.get("result") for it in m["items"]):
        out.append("<h2>Summary</h2><table><tr><th>#</th><th>Test point</th><th>Limit</th><th>Result</th></tr>")
        for i, it in enumerate(m["items"], 1):
            if it.get("result"):
                out.append(f"<tr><td>{i}</td><td>{H.escape(it.get('label') or '')}</td>"
                           f"<td>{H.escape(it.get('limit') or '')}</td><td>{badge(it['result'])}</td></tr>")
        out.append("</table>")
    out.append("<h2>Results</h2>")
    for i, it in enumerate(m["items"], 1):
        out.append(f'<div class="item"><h3>{i}. {H.escape(it.get("label") or it["type"].title())} '
                   f'{badge(it.get("result"))}</h3>')
        if it.get("limit"):
            out.append(f"<p><b>Limit:</b> {H.escape(it['limit'])}</p>")
        if it.get("note"):
            out.append(f"<p>{H.escape(it['note'])}</p>")
        img = it.get("file") if it["type"] == "screenshot" else it.get("plot")
        if img:
            out.append(f'<img src="{img64(img)}">')
        if it.get("settings"):
            out.append(f'<div class="set">{H.escape(_settings_line(it["settings"]))}</div>')
        if it.get("measurements"):
            out.append("<table><tr><th>Source</th><th>Parameter</th><th>Value</th></tr>")
            for row in _meas_rows(it["measurements"]):
                out.append("<tr>" + "".join(f"<td>{H.escape(x)}</td>" for x in row) + "</tr>")
            out.append("</table>")
        if it.get("stats"):
            out.append("<table><tr><th>Source</th><th>Min</th><th>Max</th><th>Mean</th><th>RMS</th><th>Pk-Pk</th></tr>")
            for s, v in it["stats"].items():
                out.append(f"<tr><td>{s.replace('CHAN', 'CH')}</td>" + "".join(
                    f"<td>{fmt_eng(v[k], 'V')}</td>" for k in ["min_V", "max_V", "mean_V", "rms_V", "pp_V"]) + "</tr>")
            out.append(f'</table><div class="set">Data file: {H.escape(it["file"])}</div>')
        out.append("</div>")
    if m.get("conclusion"):
        out.append(f"<h2>Conclusion</h2><p>{H.escape(m['conclusion'])}</p>")
    out.append("</body></html>")
    path.write_text("\n".join(out), encoding="utf-8")


@mcp.tool
def report_export(fmt: Annotated[str, Field(description="docx | html | both")] = "both") -> dict:
    """Generate the report document(s) from the active session."""
    if not report.active:
        raise RuntimeError("No active report.")
    base = report.dir / re.sub(r"[^\w\-]+", "_", report.meta.get("title", "report"), flags=re.UNICODE)
    files = []
    if fmt in ("docx", "both"):
        _export_docx(base.with_suffix(".docx"))
        files.append(str(base.with_suffix(".docx")))
    if fmt in ("html", "both"):
        _export_html(base.with_suffix(".html"))
        files.append(str(base.with_suffix(".html")))
    return {"files": files, "items": len(report.meta["items"])}


if __name__ == "__main__":
    mcp.run()
